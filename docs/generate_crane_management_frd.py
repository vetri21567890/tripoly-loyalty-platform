from __future__ import annotations

from datetime import date
from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Crane_Management_System_High_Level_FRD.docx"
FLOWCHART = ROOT / "crane_management_high_level_flow.png"

INK = RGBColor(17, 24, 39)
BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(88, 99, 115)
GRAY_FILL = "F3F6FA"
BLUE_FILL = "E8F1FB"
GREEN_FILL = "EAF7EE"
AMBER_FILL = "FFF6DF"
BORDER = "B9C4D0"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: float = 9.5) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], header_fill: str = GRAY_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        style_cell_text(cell, header, bold=True, color=BLUE, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            set_cell_border(cells[idx], color="D6DEE8", size="4")
            set_cell_margins(cells[idx])
            style_cell_text(cells[idx], value, size=9.25)
    doc.add_paragraph()
    return table


def add_callout(doc: Document, title: str, body: str, fill: str = BLUE_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="A9C2DD")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Aptos"
    title_run.font.size = Pt(10)
    title_run.font.color.rgb = BLUE
    paragraph.add_run("\n")
    body_run = paragraph.add_run(body)
    body_run.font.name = "Aptos"
    body_run.font.size = Pt(9.5)
    body_run.font.color.rgb = INK
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)


def get_font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/aptos.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_wrapped_text(draw: ImageDraw.ImageDraw, box, text: str, font, fill, align="center") -> None:
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 26
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    line_height = font.size + 5
    total_height = len(lines) * line_height
    y = y1 + ((y2 - y1) - total_height) / 2
    for current in lines:
        bbox = draw.textbbox((0, 0), current, font=font)
        text_width = bbox[2] - bbox[0]
        x = x1 + ((x2 - x1) - text_width) / 2 if align == "center" else x1 + 14
        draw.text((x, y), current, font=font, fill=fill)
        y += line_height


def create_flowchart() -> None:
    image = Image.new("RGB", (1600, 1120), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(38, bold=True)
    label_font = get_font(23, bold=True)
    small_font = get_font(19)
    draw.text((70, 42), "Crane Management System - High-Level Functional Flow", font=title_font, fill=(20, 40, 65))

    boxes = [
        ("Zoho CRM Sales Order", "CRM trigger starts integration", (80, 140, 390, 250), (232, 241, 251)),
        ("Customer Validation", "Reuse existing; create new in Zoho Creator", (475, 140, 785, 250), (245, 248, 252)),
        ("Crane Master", "Operations registers crane; status defaults to Idle", (870, 140, 1180, 250), (234, 247, 238)),
        ("Specification Sheet", "Auto-populate master; enter or import specs", (1230, 140, 1540, 250), (234, 247, 238)),
        ("Customer Assignment", "Assign customer and project", (80, 360, 390, 470), (255, 246, 223)),
        ("Packing List", "Create from specs; complete validation", (475, 360, 785, 470), (255, 246, 223)),
        ("Dispatch", "Create and complete dispatch record", (870, 360, 1180, 470), (232, 241, 251)),
        ("Deployment", "Capture and complete deployment", (1230, 360, 1540, 470), (232, 241, 251)),
        ("Work Started", "Crane active at customer site", (80, 590, 390, 700), (234, 247, 238)),
        ("Work Completed", "Operations marks work complete", (475, 590, 785, 700), (234, 247, 238)),
        ("Reassignment Decision", "Existing customer or new customer", (870, 590, 1180, 700), (245, 248, 252)),
        ("Next Deployment Cycle", "Crane returns to assignment workflow", (1230, 590, 1540, 700), (232, 241, 251)),
    ]

    for title, subtitle, coords, fill in boxes:
        draw.rounded_rectangle(coords, radius=18, fill=fill, outline=(137, 154, 174), width=3)
        x1, y1, x2, y2 = coords
        draw_wrapped_text(draw, (x1 + 10, y1 + 10, x2 - 10, y1 + 62), title, label_font, (20, 40, 65))
        draw_wrapped_text(draw, (x1 + 14, y1 + 64, x2 - 14, y2 - 10), subtitle, small_font, (67, 82, 101))

    def arrow(start, end):
        draw.line([start, end], fill=(31, 78, 121), width=5)
        sx, sy = start
        ex, ey = end
        if abs(ex - sx) >= abs(ey - sy):
            direction = 1 if ex > sx else -1
            points = [(ex, ey), (ex - 18 * direction, ey - 10), (ex - 18 * direction, ey + 10)]
        else:
            direction = 1 if ey > sy else -1
            points = [(ex, ey), (ex - 10, ey - 18 * direction), (ex + 10, ey - 18 * direction)]
        draw.polygon(points, fill=(31, 78, 121))

    arrow((390, 195), (475, 195))
    arrow((785, 195), (870, 195))
    arrow((1180, 195), (1230, 195))
    arrow((1540, 250), (1540, 305))
    arrow((1540, 305), (235, 305))
    arrow((235, 305), (235, 360))
    arrow((390, 415), (475, 415))
    arrow((785, 415), (870, 415))
    arrow((1180, 415), (1230, 415))
    arrow((1540, 470), (1540, 530))
    arrow((1540, 530), (235, 530))
    arrow((235, 530), (235, 590))
    arrow((390, 645), (475, 645))
    arrow((785, 645), (870, 645))
    arrow((1180, 645), (1230, 645))
    arrow((1385, 700), (1385, 820))
    arrow((1385, 820), (235, 820))
    arrow((235, 820), (235, 470))

    draw.rounded_rectangle((80, 900, 1540, 1045), radius=16, fill=(248, 250, 252), outline=(203, 213, 225), width=2)
    note = (
        "Supporting maintenance flow: HMR threshold monitoring -> email alert to Operations -> Job Card creation -> "
        "Spare Parts creation -> Inventory linkage and stock monitoring."
    )
    draw_wrapped_text(draw, (120, 922, 1500, 1025), note, get_font(24), (31, 41, 55), align="left")
    image.save(FLOWCHART)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, RGBColor(39, 72, 105), 8, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.add_run("High-Level Functional Requirement Document")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = MUTED

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("Crane Management System | Zoho Creator + Zoho CRM")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = MUTED


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Crane Management System")
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = BLUE

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle_run = subtitle.add_run("High-Level Functional Requirement Document (FRD)")
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(15)
    subtitle_run.font.color.rgb = INK

    add_callout(
        doc,
        "Purpose",
        "This document defines the high-level functional requirements and business process flows for a Crane Management System built on Zoho Creator and integrated with Zoho CRM.",
    )

    add_table(
        doc,
        ["Document Attribute", "Details"],
        [
            ["Prepared For", "Business, Operations, Sales, and Implementation Teams"],
            ["Platform", "Zoho Creator"],
            ["Primary Integration", "Zoho CRM Sales Order and Customer data"],
            ["Document Version", "1.0"],
            ["Date", date.today().strftime("%d %b %Y")],
            ["Status", "Draft for stakeholder review"],
        ],
        [2.1, 4.2],
        header_fill=BLUE_FILL,
    )
    doc.add_page_break()


def build_doc() -> None:
    create_flowchart()
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_title_page(doc)

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "The Crane Management System will be implemented in Zoho Creator to manage customer-linked crane operations, crane master data, specification sheets, packing lists, dispatch, deployment, work completion, maintenance job cards, spare parts, and operational status tracking. The application will integrate with Zoho CRM so that Sales Order creation can trigger customer validation and customer creation in Zoho Creator."
    )
    doc.add_paragraph(
        "The solution is intended to provide a guided operations workflow from initial customer assignment through deployment and reassignment, while preserving visibility into crane availability, maintenance activity, specification details, spare parts usage, and inventory linkages."
    )

    doc.add_heading("2. Business Objectives", level=1)
    add_bullets(
        doc,
        [
            "Automate customer creation in Zoho Creator based on Zoho CRM Sales Order activity.",
            "Maintain a centralized Crane Master with clear operational status tracking.",
            "Reduce manual entry by auto-populating crane details and importing specification data through a Zoho Sheet template.",
            "Guide crane deployment through controlled workflow stages and stage-based validations.",
            "Trigger maintenance actions based on HMR thresholds and support Job Card creation.",
            "Link spare parts consumption with Job Cards and Inventory for stock visibility.",
            "Support repeated crane deployment cycles for existing or new customers.",
        ],
    )

    doc.add_heading("3. Scope", level=1)
    add_table(
        doc,
        ["In Scope", "Out of Scope / Pending"],
        [
            ["Zoho CRM Sales Order trigger and customer validation", "Detailed rental billing logic"],
            ["Customer, Crane, Specification Sheet, Packing List, Dispatch, Deployment, Job Card, Spare Parts, and Inventory-linked workflows", "Billing approvals and finance posting rules"],
            ["Crane blueprint-style deployment workflow and status transitions", "Detailed log journal and audit reporting rules"],
            ["HMR-based maintenance alert and Job Card initiation", "Advanced analytics dashboards, unless separately confirmed"],
        ],
        [3.15, 3.15],
        header_fill=GRAY_FILL,
    )

    doc.add_heading("4. Users and Roles", level=1)
    add_table(
        doc,
        ["Role", "Primary Responsibilities"],
        [
            ["Sales / CRM Users", "Create Sales Orders in Zoho CRM and maintain customer-facing commercial information."],
            ["Operations Team", "Create cranes, assign customers, manage packing lists, dispatch, deployment, work completion, reassignment, and maintenance actions."],
            ["Inventory Team", "Maintain inventory items and monitor spare parts stock levels."],
            ["System Administrator", "Configure integrations, workflows, permissions, import templates, validations, and notifications."],
            ["Management / Reviewers", "Monitor operational status, crane utilization, maintenance activity, and pending actions."],
        ],
        [1.75, 4.55],
    )

    doc.add_heading("5. High-Level Functional Process Flow", level=1)
    doc.add_paragraph(
        "The following diagram represents the overall structure-level process flow from Zoho CRM Sales Order creation through customer validation, crane setup, deployment workflow, work completion, and reassignment."
    )
    doc.add_picture(str(FLOWCHART), width=Inches(6.45))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("6. Functional Requirements", level=1)

    doc.add_heading("6.1 Customer Integration from Zoho CRM", level=2)
    add_numbered(
        doc,
        [
            "Once a Sales Order is created in Zoho CRM, the integration process shall be triggered.",
            "The system shall check whether the customer already exists in the Crane Management System in Zoho Creator.",
            "If the customer exists, no duplicate Customer record shall be created.",
            "If the customer does not exist, the system shall automatically create a new Customer record in the Customer module.",
        ],
    )
    add_table(
        doc,
        ["Requirement ID", "Requirement"],
        [
            ["FR-CRM-01", "Trigger customer synchronization when a Sales Order is created in Zoho CRM."],
            ["FR-CRM-02", "Validate customer existence using a unique customer identifier, such as CRM customer ID, customer code, email, phone, or configured matching logic."],
            ["FR-CRM-03", "Create a new Customer record in Zoho Creator only when the customer does not already exist."],
            ["FR-CRM-04", "Maintain CRM reference fields in Zoho Creator for traceability."],
        ],
        [1.25, 5.05],
    )

    doc.add_heading("6.2 Crane Creation Process", level=2)
    add_bullets(
        doc,
        [
            "Operations users shall create Crane records in the Crane Module.",
            "Crane details shall be entered manually during crane registration.",
            "The default Crane Status shall be set to Idle upon creation.",
        ],
    )

    doc.add_heading("6.3 Crane Specification Sheet Creation", level=2)
    add_bullets(
        doc,
        [
            "Users shall create a Crane Specification Sheet by selecting the corresponding Crane record.",
            "Crane Master details shall automatically populate into the Specification Sheet after crane selection.",
            "Users shall manually enter detailed specification line items in the Specification Sheet Subform.",
            "The system shall provide a Zoho Sheet template for bulk upload of specification details.",
            "When the user clicks Import, the uploaded specification details shall populate into the Specification Sheet Subform.",
        ],
    )

    doc.add_heading("6.4 Job Card Management", level=2)
    add_bullets(
        doc,
        [
            "The system shall monitor Crane HMR values against predefined maintenance thresholds.",
            "When an HMR threshold is reached, an automated email notification shall be sent to the Operations Team.",
            "A Job Card creation button shall be available within the Crane record.",
            "Clicking the button shall open the Job Card creation form, auto-link the Crane record, and allow users to enter maintenance details.",
        ],
    )

    doc.add_heading("6.5 Spare Parts Management", level=2)
    add_bullets(
        doc,
        [
            "A Create Spare Parts button shall be available inside the Job Card.",
            "Clicking the button shall open the Spare Parts form and link the spare part record with the respective Job Card.",
            "Spare Parts records shall link inventory items from the Inventory Module.",
            "Inventory stock levels shall be monitored and maintained through the inventory integration.",
        ],
    )

    doc.add_heading("6.6 Crane Deployment Workflow", level=2)
    doc.add_paragraph(
        "After crane creation, the crane shall move through operational stages using guided workflow transitions. Each stage will capture required data, apply validations, and update the crane status as applicable."
    )
    add_table(
        doc,
        ["Stage", "Activities", "Validation / Output"],
        [
            ["1. Customer Assignment", "Assign customer and enter project-specific deployment details.", "User clicks Move to Next Stage. Crane status becomes Customer Assigned."],
            ["2. Packing List Creation", "Create Packing List and associate Crane Specification Sheet.", "Specification Sheet subform items auto-populate into Packing List subform. Users can modify items."],
            ["3. Packing List Completion", "User clicks Packing List Done.", "System verifies Packing List exists and status is Completed. If valid, Crane Status becomes Packing List Completed."],
            ["4. Dispatch Creation", "Create Dispatch record and enter dispatch details.", "Dispatch must be completed before proceeding to Deployment."],
            ["5. Deployment", "Create Deployment record and capture deployment details.", "Deployment status must be completed before moving to Work Started."],
            ["6. Work Started", "Crane is actively deployed at customer site and operational work begins.", "Status becomes Work In Progress."],
            ["7. Work Completed", "Operations marks work as completed.", "System prompts for the next assignment."],
            ["8. Customer Reassignment", "User selects Existing Customer or New Customer.", "Existing customer restarts deployment workflow; new customer popup creates a Customer record and starts next cycle."],
        ],
        [1.55, 2.6, 2.15],
        header_fill=BLUE_FILL,
    )

    doc.add_heading("7. Status Flow Summary", level=1)
    add_callout(
        doc,
        "Crane Status Lifecycle",
        "Idle -> Customer Assigned -> Packing List Created -> Packing List Completed -> Dispatch Created -> Deployment Completed -> Work Started -> Work Completed -> Customer Reassignment -> Next Deployment Cycle",
        fill=GREEN_FILL,
    )

    doc.add_heading("8. Integration Requirements", level=1)
    add_table(
        doc,
        ["Integration Area", "Expected Behavior"],
        [
            ["Zoho CRM to Zoho Creator", "Sales Order creation triggers customer validation and optional customer creation in the Crane Management System."],
            ["Zoho Sheet Import", "Approved template supports bulk import of specification details into the Specification Sheet Subform."],
            ["Inventory Module", "Spare Parts records link to inventory items so stock can be monitored and maintained."],
            ["Email Notification", "HMR threshold alerts notify the Operations Team and prompt Job Card creation."],
        ],
        [2.0, 4.3],
    )

    doc.add_heading("9. Key Validations and Controls", level=1)
    add_bullets(
        doc,
        [
            "Customer records shall not be duplicated during CRM integration.",
            "Crane Status shall default to Idle during Crane creation.",
            "Packing List Completion shall require an existing Packing List with status Completed.",
            "Dispatch shall be completed before moving to Deployment.",
            "Deployment shall be completed before moving to Work Started.",
            "Customer reassignment shall require the user to select whether the next customer is existing or new.",
            "System workflows shall show clear error messages when validation fails and prevent invalid stage transitions.",
        ],
    )

    doc.add_heading("10. Reporting and Visibility", level=1)
    add_bullets(
        doc,
        [
            "Crane status report by stage and customer assignment.",
            "Crane availability report showing Idle, Assigned, In Progress, and Completed statuses.",
            "Maintenance alert and Job Card status report.",
            "Spare parts usage and inventory linkage report.",
            "Packing List, Dispatch, and Deployment completion tracking.",
        ],
    )

    doc.add_heading("11. Assumptions and Dependencies", level=1)
    add_bullets(
        doc,
        [
            "Zoho CRM Sales Orders will contain sufficient customer information for validation and customer creation.",
            "A unique matching rule for customer validation will be finalized during implementation.",
            "HMR threshold values will be defined by the Operations Team.",
            "Zoho Sheet import templates will follow a controlled column structure.",
            "User roles and field-level permissions will be configured in Zoho Creator based on operational responsibilities.",
        ],
    )

    doc.add_heading("12. Pending Discussion Items", level=1)
    add_table(
        doc,
        ["Module / Process", "Pending Items"],
        [
            ["Billing Module", "Rental billing process, billing approvals, billing generation logic, and finance handoff rules."],
            ["Log Journal Module", "Crane activity logs, usage tracking, operational history, and audit record requirements."],
        ],
        [2.0, 4.3],
        header_fill=AMBER_FILL,
    )

    doc.add_heading("13. Acceptance Criteria", level=1)
    add_bullets(
        doc,
        [
            "Sales Order creation in Zoho CRM triggers customer validation in Zoho Creator.",
            "Duplicate customer records are prevented based on the agreed matching logic.",
            "Crane creation defaults the Crane Status to Idle.",
            "Specification Sheet creation auto-populates Crane Master details after crane selection.",
            "Specification details can be populated through manual entry or bulk import.",
            "Packing List items can be generated from the selected Specification Sheet and modified by users.",
            "Workflow validations prevent users from skipping required completion steps.",
            "HMR threshold alerts are sent to Operations and support Job Card creation.",
            "Spare Parts records are linked to Job Cards and Inventory items.",
            "Completed work can move into customer reassignment and a new deployment cycle.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
