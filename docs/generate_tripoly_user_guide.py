from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Tripoly_Loyalty_Platform_User_Guide_and_Structure.docx"


INK = RGBColor(15, 35, 29)
GREEN = RGBColor(22, 163, 74)
MUTED = RGBColor(82, 92, 108)
LIGHT_GREEN = "ECFDF3"
LIGHT_GRAY = "F8FAFC"
BORDER = "CBD5E1"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GREEN)
        set_cell_text(cell, header, bold=True, color=INK)
        if widths:
            cell.width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            if widths:
                cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()
    return table


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_steps(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREEN)
    p = cell.paragraphs[0]
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = INK
    p.add_run("\n")
    body_run = p.add_run(body)
    body_run.font.name = "Aptos"
    body_run.font.size = Pt(9.5)
    body_run.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_paragraph()


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for line in text.splitlines():
        run = p.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)


def build_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color, before, after in [
        ("Heading 1", 18, GREEN, 16, 6),
        ("Heading 2", 14, INK, 12, 5),
        ("Heading 3", 11.5, INK, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Heading 1" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Tripoly Loyalty Platform")
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = INK
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("User Guide and Structural Sketch")
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = GREEN
    meta = doc.add_paragraph()
    meta_run = meta.add_run("Prepared for Admin and Customer portal handover | FastAPI, PostgreSQL, React/Vite, Docker Compose")
    meta_run.font.name = "Aptos"
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = MUTED

    add_callout(
        doc,
        "Document purpose",
        "This guide explains how Tripoly works as an independent travel loyalty and retention platform. "
        "It covers day-to-day Admin and Customer usage, the platform structure, core data visibility rules, and deployment/testing commands.",
    )

    doc.add_heading("1. Platform Summary", level=1)
    doc.add_paragraph(
        "Tripoly is an independent loyalty and retention platform for travel businesses. It manages Travel Coins, rewards, tiers, "
        "referrals, missions, campaigns, UGC moderation, notifications, analytics, and audit visibility. It is not an ecommerce storefront, cart, checkout, or order-management system."
    )
    add_bullets(
        doc,
        [
            "Customer portal: port 5173, used by travellers to view coins, rewards, referrals, missions, campaigns, UGC, notifications, and profile.",
            "Admin portal: port 5174, used by operators to manage customers, loyalty program configuration, campaigns, rewards, UGC, reports, and audit logs.",
            "API service: port 8000, FastAPI backend exposing /api/v1 endpoints and partner integration surfaces.",
            "Database: PostgreSQL stores customers, wallets, rewards, campaigns, tiers, admin users, roles, and audit data.",
            "External travel systems connect through APIs/webhooks only. Zoho CRM, booking systems, websites, WhatsApp tools, and partner systems are not tightly coupled.",
        ],
    )

    doc.add_heading("2. Structural Sketch", level=1)
    doc.add_paragraph("High-level runtime structure:")
    add_code_block(
        doc,
        """[Admin Browser]    [Customer Browser]
      |                  |
      v                  v
Admin Vite App      Customer Vite App
  Port 5174           Port 5173
      |                  |
      +-------- HTTPS/HTTP API calls --------+
                                             |
                                             v
                                  FastAPI Backend
                                     Port 8000
                                             |
                         +-------------------+-------------------+
                         v                                       v
                    PostgreSQL                                Redis
              loyalty data, RBAC,                         OTP/cache
              wallet ledger, audits

External systems: Zoho CRM, booking engine, website, WhatsApp, travel ops
connect only through API/webhook integration surfaces.""",
    )

    add_table(
        doc,
        ["Layer", "Technology", "Responsibility"],
        [
            ["Frontend Admin", "React 18 + Vite", "Operator workspace for loyalty and travel-retention management."],
            ["Frontend Customer", "React 18 + Vite", "Traveller self-service portal for wallet, rewards, referrals, missions, and campaigns."],
            ["Backend API", "FastAPI", "Authentication, RBAC, wallet ledger, program rules, rewards, campaigns, reports, and integrations."],
            ["Database", "PostgreSQL + SQLAlchemy async", "Durable storage for users, roles, wallets, transactions, configuration, and audit data."],
            ["Migrations", "Alembic", "Database schema changes."],
            ["Runtime", "Docker Compose", "Local and shareable development orchestration."],
        ],
        widths=[1.35, 1.7, 3.35],
    )

    doc.add_heading("3. Access and Credentials", level=1)
    add_table(
        doc,
        ["Surface", "Local URL", "Purpose"],
        [
            ["Admin portal", "http://localhost:5174", "Admin login, management, reporting, and configuration."],
            ["Customer portal", "http://localhost:5173", "Customer signup/login and loyalty self-service."],
            ["API", "http://localhost:8000", "Backend service; health, docs, and /api/v1 endpoints."],
            ["API docs", "http://localhost:8000/docs", "Swagger documentation."],
        ],
        widths=[1.4, 2.25, 2.75],
    )
    add_callout(
        doc,
        "Default admin",
        "Email: admin@tripoly.app\nPassword: admin123\nReset command: docker compose exec api python scripts/reset_admin_password.py",
    )

    doc.add_heading("4. Admin User Guide", level=1)
    doc.add_paragraph("Admin users operate the loyalty program from the Admin portal. Admin routes require an admin JWT and backend role validation.")

    doc.add_heading("4.1 Admin Login and Logout", level=2)
    add_steps(
        doc,
        [
            "Open http://localhost:5174 or the Admin ngrok URL.",
            "Enter admin@tripoly.app and admin123.",
            "After login, confirm the dark sidebar and Admin Workspace header are visible.",
            "Use Logout from the sidebar to clear local storage, session storage, auth state, and authorization headers.",
        ],
    )

    doc.add_heading("4.2 Admin Navigation Map", level=2)
    add_table(
        doc,
        ["Navigation Area", "Modules", "Admin Purpose"],
        [
            ["Core", "Dashboard, Customers", "Monitor platform health and manage customer accounts."],
            ["Loyalty Program", "Points Engine, Rewards, Tiers/VIP, Missions, Referrals, Campaigns, Wallet Rules, Notifications, Program Settings, Performance", "Configure and operate loyalty mechanics."],
            ["UGC and Reviews", "UGC Moderation, Testimonials, Review Rewards", "Review travel content and reward eligible submissions."],
            ["Travel Operations", "Booking Webhooks, Partner Integrations, Travel Activity Logs", "Inspect travel-system integration activity."],
            ["Reports", "Customer, Wallet, Referral, Campaign, UGC, Redemption Analytics", "View operational and loyalty performance data."],
            ["System", "Audit Logs, Settings", "Review admin activity and platform settings."],
        ],
        widths=[1.35, 2.55, 2.5],
    )

    doc.add_heading("4.3 Customer Management", level=2)
    add_bullets(
        doc,
        [
            "View all customer records with email, phone, name, status, available coins, lifetime coins, and joined date.",
            "Search customers by name, email, phone, or status.",
            "Add a customer from the Admin portal. The backend creates the customer account, profile row, and zero-balance wallet.",
            "Export the visible customer list as CSV for operational reporting.",
        ],
    )

    doc.add_heading("4.4 Campaign Management", level=2)
    add_bullets(
        doc,
        [
            "Create or edit travel-loyalty campaigns.",
            "Set campaign name, description, start date, end date, eligibility type, reward coins, tasks/missions, and status.",
            "Supported eligibility: ALL_CUSTOMERS, ENROLLED_ONLY, INVITE_ONLY.",
            "Supported statuses: draft, active, ended, archived.",
            "Activate or deactivate campaigns from the table action buttons.",
        ],
    )
    add_table(
        doc,
        ["Seed Campaign", "Eligibility", "Status", "Reward Coins"],
        [
            ["Japan Launch Campaign", "ALL_CUSTOMERS", "active", "750"],
            ["Summer Escape Campaign", "ENROLLED_ONLY", "active", "500"],
            ["Insider Travel Deal", "INVITE_ONLY", "active", "1000"],
            ["Festive Travel Rewards", "ALL_CUSTOMERS", "draft", "650"],
        ],
        widths=[2.25, 1.55, 1.15, 1.1],
    )

    doc.add_heading("4.5 Tiers / VIP", level=2)
    add_bullets(
        doc,
        [
            "Manage tier names, thresholds, benefits, display order, and active/inactive state.",
            "Default tiers: Explorer, Voyager, Elite Traveller, Global Nomad.",
            "Tier calculation uses lifetime coins earned, not current wallet balance.",
            "Redeeming coins does not reduce tier qualification.",
        ],
    )

    doc.add_heading("4.6 Wallet Rules", level=2)
    add_bullets(
        doc,
        [
            "Set coins_per_rupee, coin expiry months, minimum redemption coins, partial redemption allowed, and optional max redemption per booking.",
            "Default conversion: 10 Travel Coins = INR 1.",
            "Wallet redemptions use FIFO coin grants and preserve lifetime earned coins.",
        ],
    )

    doc.add_heading("4.7 Rewards and Redemptions", level=2)
    add_bullets(
        doc,
        [
            "Create and manage rewards such as travel vouchers, room upgrades, lounge passes, discounts, experiences, cashback vouchers, and merchandise.",
            "Customers can redeem only when they have sufficient available coins.",
            "Reward redemption generates voucher history and deducts coins through the wallet ledger.",
        ],
    )

    doc.add_heading("4.8 UGC Moderation", level=2)
    add_bullets(
        doc,
        [
            "Review pending customer content submissions.",
            "Approve eligible UGC and issue configured coin rewards.",
            "Reject submissions with a rejection reason so customers can see status and resubmit where applicable.",
        ],
    )

    doc.add_heading("4.9 Reports and Audit Logs", level=2)
    add_bullets(
        doc,
        [
            "Reports are view-only unless a real action exists.",
            "Audit logs provide administrative visibility for operational control.",
            "Customers cannot access reports, audit logs, settings, or admin configuration.",
        ],
    )

    doc.add_heading("5. Customer User Guide", level=1)
    doc.add_paragraph("Customers use the Customer portal to manage their own loyalty activity. Customer routes require customer JWTs and return only that customer's data.")

    doc.add_heading("5.1 Customer Login and Signup", level=2)
    add_steps(
        doc,
        [
            "Open http://localhost:5173 or the Customer ngrok URL.",
            "Sign up with email and password, or sign in with an existing account.",
            "After login, confirm the sidebar shows Dashboard, Wallet, Rewards, Referrals, UGC Uploads, Missions, Campaigns, Profile, and Notifications.",
            "Use Logout from the sidebar to clear customer auth state and API authorization headers.",
        ],
    )

    doc.add_heading("5.2 Dashboard", level=2)
    add_bullets(
        doc,
        [
            "Shows Travel Coins balance and approximate INR redemption value.",
            "Shows tier, referral code, conversion rule, expiring coins, recent transactions, rewards, campaigns, and missions.",
            "Highlights active campaigns and mission progress.",
        ],
    )

    doc.add_heading("5.3 Wallet", level=2)
    add_bullets(
        doc,
        [
            "Shows available coins, lifetime coins, redeemed coins, expired coins, and approximate INR value.",
            "Shows the conversion rule: 10 Travel Coins = INR 1 by default.",
            "Lists transaction history and expiring-soon coin grants.",
        ],
    )

    doc.add_heading("5.4 Rewards", level=2)
    add_bullets(
        doc,
        [
            "Shows available rewards, coin cost, INR equivalent, inventory status, and customer balance.",
            "Prevents redemption when customer coins are insufficient.",
            "Stores redemption history and voucher code after successful redemption.",
        ],
    )

    doc.add_heading("5.5 Referrals", level=2)
    add_bullets(
        doc,
        [
            "Shows referral code and referral link.",
            "Allows copying the referral link.",
            "Tracks pending referrals, qualified referrals, and reward events.",
        ],
    )

    doc.add_heading("5.6 Missions / Tasks", level=2)
    add_bullets(
        doc,
        [
            "Shows available missions such as complete profile, upload review, upload photo, upload video/reel, refer friend, and campaign tasks.",
            "Shows mission progress and status.",
            "Coins are awarded according to configured mission and approval rules.",
        ],
    )

    doc.add_heading("5.7 Campaigns", level=2)
    add_bullets(
        doc,
        [
            "Shows active eligible campaigns only.",
            "ALL_CUSTOMERS campaigns are visible to all active customers.",
            "ENROLLED_ONLY campaigns show a Join Campaign button until enrolled.",
            "INVITE_ONLY campaigns remain hidden unless the customer has an invitation.",
        ],
    )

    doc.add_heading("5.8 UGC Uploads", level=2)
    add_bullets(
        doc,
        [
            "Submit travel photos, review screenshots, testimonials, videos, or reels.",
            "Submissions enter admin moderation.",
            "Approved submissions can earn configured coin rewards.",
            "Rejected submissions display the rejection reason.",
        ],
    )

    doc.add_heading("5.9 Notifications and Profile", level=2)
    add_bullets(
        doc,
        [
            "Notifications include reward, voucher, tier, campaign, referral, and UGC approval/rejection updates.",
            "Profile lets customers update first and last name.",
            "Completing profile can award mission coins when configured.",
        ],
    )

    doc.add_heading("6. Visibility and RBAC Rules", level=1)
    add_table(
        doc,
        ["Actor", "Can Access", "Must Not Access"],
        [
            ["Customer", "Own profile, wallet, transactions, referrals, UGC, missions, campaigns, notifications, eligible rewards.", "Other customer data, global analytics, audit logs, admin settings, earning rules, API keys, webhook secrets."],
            ["Admin", "All customers, loyalty configuration, rewards, campaigns, referrals, UGC queue, wallet activity, reports, audit logs.", "Customer passwords or raw secrets. Admin should operate through role-protected APIs."],
        ],
        widths=[1.0, 2.85, 2.85],
    )

    doc.add_heading("7. Public Sharing and Local Development", level=1)
    add_bullets(
        doc,
        [
            "Local access: Admin http://localhost:5174, Customer http://localhost:5173, API http://localhost:8000.",
            "Public sharing requires separate ngrok tunnels for API, Admin, and Customer.",
            "Frontend VITE_API_BASE_URL must point to the API ngrok URL with /api/v1 appended.",
            "Backend CORS allows localhost, LAN/private IPs, ngrok-free.app, and ngrok.io via configuration.",
            "For free ngrok tunnels, frontend axios sends ngrok-skip-browser-warning: true so API calls do not receive the ngrok warning page.",
        ],
    )
    add_code_block(
        doc,
        """# Local mode
VITE_API_BASE_URL=http://localhost:8000/api/v1

# Public sharing mode
VITE_API_BASE_URL=https://YOUR_API_NGROK_URL/api/v1

# Start/rebuild
docker compose down
docker compose up -d --build

# Ensure admin credentials
docker compose exec api python scripts/reset_admin_password.py""",
    )

    doc.add_heading("8. Operational Commands", level=1)
    add_code_block(
        doc,
        """docker compose down
docker compose up -d --build
docker compose exec api alembic upgrade head
docker compose exec api python scripts/seed.py
docker compose exec api python scripts/reset_admin_password.py
docker compose logs api --tail=200
docker compose exec admin npm run build
docker compose exec web npm run build""",
    )

    doc.add_heading("9. Testing Checklist", level=1)
    add_table(
        doc,
        ["Area", "Check"],
        [
            ["API", "GET /health returns {'status':'ok'} and /docs loads."],
            ["Admin auth", "Login with admin@tripoly.app / admin123 succeeds locally and through ngrok."],
            ["Admin logout", "Logout clears auth state and redirects to /login."],
            ["Customer auth", "Signup/signin works; customer logout clears auth state."],
            ["Campaigns", "Admin can create/edit/activate/deactivate; customer sees only eligible active campaigns."],
            ["Wallet", "Customer sees coin balance, INR equivalent, conversion rule, transactions, expiring coins."],
            ["Rewards", "Insufficient coins block redemption; successful redemption creates voucher history."],
            ["RBAC", "Customer token cannot access admin endpoints; admin endpoints require admin role."],
            ["Public share", "Admin and customer ngrok URLs call API ngrok URL, not viewer localhost."],
        ],
        widths=[1.45, 5.1],
    )

    doc.add_heading("10. Handover Notes", level=1)
    add_bullets(
        doc,
        [
            "Tripoly should stay independent from ecommerce flows. Do not add product cart, checkout, storefront, marketplace, Shopify/Wix, or order-management logic.",
            "Use APIs/webhooks for external travel operations systems.",
            "Keep role checks in the backend. Frontend route guards are useful but never sufficient for security.",
            "Keep Travel Coins conversion centralized through program settings rather than hardcoding it in pages.",
            "When recreating containers or databases, always run migrations and seed/reset scripts before testing login.",
        ],
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Tripoly Loyalty Platform - User Guide and Structural Sketch")
    fr.font.name = "Aptos"
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
